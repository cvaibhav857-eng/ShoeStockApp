from flask import Flask, render_template, request, redirect, send_file
import sqlite3
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


app = Flask(__name__)


DATABASE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "stock.db"
)


# =========================================================
# DATABASE CONNECTION
# =========================================================

def get_db():

    conn = sqlite3.connect(
        DATABASE,
        timeout=10
    )

    conn.execute(
        "PRAGMA busy_timeout = 10000"
    )

    return conn


# =========================================================
# DATABASE TABLES
# =========================================================

def init_db():

    conn = get_db()
    cursor = conn.cursor()

    # -----------------------------------------------------
    # ARTICLES
    # -----------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS articles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            article_code TEXT UNIQUE,
            article_name TEXT,
            size_2 INTEGER DEFAULT 0,
            size_3 INTEGER DEFAULT 0,
            size_4 INTEGER DEFAULT 0,
            size_5 INTEGER DEFAULT 0,
            size_6 INTEGER DEFAULT 0,
            size_7 INTEGER DEFAULT 0,
            size_8 INTEGER DEFAULT 0,
            size_9 INTEGER DEFAULT 0,
            size_10 INTEGER DEFAULT 0,
            size_11 INTEGER DEFAULT 0,
            size_12 INTEGER DEFAULT 0,
            size_13 INTEGER DEFAULT 0,
            size_14 INTEGER DEFAULT 0
        )
    """)

    # -----------------------------------------------------
    # DISPATCH HISTORY
    # -----------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS dispatches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            dispatch_date TEXT,
            party_name TEXT,
            article_code TEXT,
            article_name TEXT,
            size_2 INTEGER DEFAULT 0,
            size_3 INTEGER DEFAULT 0,
            size_4 INTEGER DEFAULT 0,
            size_5 INTEGER DEFAULT 0,
            size_6 INTEGER DEFAULT 0,
            size_7 INTEGER DEFAULT 0,
            size_8 INTEGER DEFAULT 0,
            size_9 INTEGER DEFAULT 0,
            size_10 INTEGER DEFAULT 0,
            size_11 INTEGER DEFAULT 0,
            size_12 INTEGER DEFAULT 0,
            size_13 INTEGER DEFAULT 0,
            size_14 INTEGER DEFAULT 0,
            total_qty INTEGER DEFAULT 0
        )
    """)

    # -----------------------------------------------------
    # STOCK IN HISTORY
    # -----------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS stock_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            stock_date TEXT,
            article_code TEXT,
            article_name TEXT,
            size_2 INTEGER DEFAULT 0,
            size_3 INTEGER DEFAULT 0,
            size_4 INTEGER DEFAULT 0,
            size_5 INTEGER DEFAULT 0,
            size_6 INTEGER DEFAULT 0,
            size_7 INTEGER DEFAULT 0,
            size_8 INTEGER DEFAULT 0,
            size_9 INTEGER DEFAULT 0,
            size_10 INTEGER DEFAULT 0,
            size_11 INTEGER DEFAULT 0,
            size_12 INTEGER DEFAULT 0,
            size_13 INTEGER DEFAULT 0,
            size_14 INTEGER DEFAULT 0,
            total_qty INTEGER DEFAULT 0
        )
    """)

    conn.commit()
    conn.close()


# =========================================================
# DASHBOARD
# =========================================================

@app.route("/")
def home():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT COUNT(*)
        FROM articles
    """)

    total_articles = cursor.fetchone()[0]

    cursor.execute("""
        SELECT COALESCE(SUM(
            size_2 + size_3 + size_4 + size_5 +
            size_6 + size_7 + size_8 + size_9 +
            size_10 + size_11 + size_12 + size_13 + size_14
        ), 0)
        FROM articles
    """)

    remaining_stock = cursor.fetchone()[0]

    cursor.execute("""
        SELECT COALESCE(SUM(total_qty), 0)
        FROM dispatches
    """)

    total_dispatch = cursor.fetchone()[0]

    total_stock = remaining_stock + total_dispatch

    conn.close()

    return render_template(
        "dashboard.html",
        total_articles=total_articles,
        total_stock=total_stock,
        total_dispatch=total_dispatch,
        remaining_stock=remaining_stock
    )


# =========================================================
# ADD ARTICLE
# =========================================================

@app.route(
    "/add-article",
    methods=["GET", "POST"]
)
def add_article():

    if request.method == "POST":

        article_code = request.form.get(
            "article_code",
            ""
        ).strip()

        article_name = request.form.get(
            "article_name",
            ""
        ).strip()

        if not article_code:

            return "Article Code enter karo."

        if not article_name:

            return "Article Name enter karo."

        sizes = {}

        for size in range(2, 15):

            value = request.form.get(
                f"size_{size}",
                "0"
            ).strip()

            try:
                quantity = int(value)
            except ValueError:
                quantity = 0

            if quantity < 0:
                quantity = 0

            sizes[size] = quantity

        conn = None

        try:

            conn = get_db()
            cursor = conn.cursor()

            cursor.execute("""
                INSERT INTO articles (
                    article_code,
                    article_name,
                    size_2,
                    size_3,
                    size_4,
                    size_5,
                    size_6,
                    size_7,
                    size_8,
                    size_9,
                    size_10,
                    size_11,
                    size_12,
                    size_13,
                    size_14
                )
                VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?, ?
                )
            """, (
                article_code,
                article_name,
                sizes[2],
                sizes[3],
                sizes[4],
                sizes[5],
                sizes[6],
                sizes[7],
                sizes[8],
                sizes[9],
                sizes[10],
                sizes[11],
                sizes[12],
                sizes[13],
                sizes[14]
            ))

            conn.commit()

            return redirect("/")

        except sqlite3.IntegrityError:

            return "Article Code already exists."

        except sqlite3.OperationalError as e:

            return f"Database Error: {e}"

        finally:

            if conn:
                conn.close()

    return render_template(
        "add_article.html"
    )


# =========================================================
# ARTICLE MANAGEMENT
# =========================================================

@app.route("/article-management")
def article_management():

    search = request.args.get(
        "search",
        ""
    ).strip()

    conn = get_db()
    cursor = conn.cursor()

    query = """
        SELECT
            id,
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14
        FROM articles
        WHERE 1=1
    """

    params = []

    if search:

        query += """
            AND (
                article_code LIKE ?
                OR article_name LIKE ?
            )
        """

        search_value = f"%{search}%"

        params.append(search_value)
        params.append(search_value)

    query += """
        ORDER BY article_code
    """

    cursor.execute(
        query,
        params
    )

    article_rows = cursor.fetchall()

    # -----------------------------------------------------
    # ADD TOTAL STOCK TO EACH ROW
    # -----------------------------------------------------

    articles = []

    for row in article_rows:

        total_stock = sum(
            row[3:16]
        )

        articles.append(
            tuple(row) + (total_stock,)
        )

    conn.close()

    return render_template(
        "article_management.html",
        articles=articles,
        search=search
    )


# =========================================================
# EDIT ARTICLE
# =========================================================

@app.route(
    "/edit-article/<int:article_id>",
    methods=["GET", "POST"]
)
def edit_article(article_id):

    conn = get_db()
    cursor = conn.cursor()

    # -----------------------------------------------------
    # GET ARTICLE
    # -----------------------------------------------------

    cursor.execute("""
        SELECT
            id,
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14
        FROM articles
        WHERE id = ?
    """, (article_id,))

    article = cursor.fetchone()

    if not article:

        conn.close()

        return "Article not found."

    # -----------------------------------------------------
    # POST - UPDATE ARTICLE
    # -----------------------------------------------------

    if request.method == "POST":

        new_code = request.form.get(
            "article_code",
            ""
        ).strip()

        new_name = request.form.get(
            "article_name",
            ""
        ).strip()

        if not new_code:

            conn.close()

            return "Article Code enter karo."

        if not new_name:

            conn.close()

            return "Article Name enter karo."

        old_code = article[1]

        # -------------------------------------------------
        # CHECK DUPLICATE ARTICLE CODE
        # -------------------------------------------------

        cursor.execute("""
            SELECT id
            FROM articles
            WHERE article_code = ?
              AND id != ?
        """, (
            new_code,
            article_id
        ))

        duplicate = cursor.fetchone()

        if duplicate:

            conn.close()

            return "Article Code already exists."

        # -------------------------------------------------
        # UPDATE ARTICLE
        # -------------------------------------------------

        try:

            cursor.execute("""
                UPDATE articles

                SET
                    article_code = ?,
                    article_name = ?

                WHERE id = ?
            """, (
                new_code,
                new_name,
                article_id
            ))

            # -------------------------------------------------
            # UPDATE STOCK IN HISTORY
            # -------------------------------------------------

            cursor.execute("""
                UPDATE stock_entries

                SET
                    article_code = ?,
                    article_name = ?

                WHERE article_code = ?
            """, (
                new_code,
                new_name,
                old_code
            ))

            # -------------------------------------------------
            # UPDATE DISPATCH HISTORY
            # -------------------------------------------------

            cursor.execute("""
                UPDATE dispatches

                SET
                    article_code = ?,
                    article_name = ?

                WHERE article_code = ?
            """, (
                new_code,
                new_name,
                old_code
            ))

            conn.commit()

            return redirect(
                "/article-management"
            )

        except sqlite3.IntegrityError:

            conn.rollback()

            return "Article Code already exists."

        except Exception as e:

            conn.rollback()

            return f"Article Update Error: {e}"

        finally:

            conn.close()

    # -----------------------------------------------------
    # GET ARTICLE EDIT PAGE
    # -----------------------------------------------------

    conn.close()

    return render_template(
        "edit_article.html",
        article=article
    )


# =========================================================
# DELETE ARTICLE
# =========================================================

@app.route(
    "/delete-article/<int:article_id>",
    methods=["POST"]
)
def delete_article(article_id):

    conn = get_db()
    cursor = conn.cursor()

    try:

        # -------------------------------------------------
        # GET ARTICLE
        # -------------------------------------------------

        cursor.execute("""
            SELECT
                article_code,
                article_name,
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14
            FROM articles
            WHERE id = ?
        """, (article_id,))

        article = cursor.fetchone()

        if not article:

            return "Article not found."

        article_code = article[0]

        # -------------------------------------------------
        # CHECK CURRENT STOCK
        # -------------------------------------------------

        current_stock = sum(
            article[2:15]
        )

        if current_stock > 0:

            return (
                "Article delete nahi ho sakta. "
                f"Current stock {current_stock} pairs hai. "
                "Pehle stock zero karo."
            )

        # -------------------------------------------------
        # CHECK STOCK IN HISTORY
        # -------------------------------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM stock_entries
            WHERE article_code = ?
        """, (article_code,))

        stock_history_count = cursor.fetchone()[0]

        if stock_history_count > 0:

            return (
                "Article delete nahi ho sakta. "
                "Is article ki Stock In History available hai."
            )

        # -------------------------------------------------
        # CHECK DISPATCH HISTORY
        # -------------------------------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM dispatches
            WHERE article_code = ?
        """, (article_code,))

        dispatch_history_count = cursor.fetchone()[0]

        if dispatch_history_count > 0:

            return (
                "Article delete nahi ho sakta. "
                "Is article ki Dispatch History available hai."
            )

        # -------------------------------------------------
        # DELETE ARTICLE
        # -------------------------------------------------

        cursor.execute("""
            DELETE FROM articles
            WHERE id = ?
        """, (article_id,))

        conn.commit()

        return redirect(
            "/article-management"
        )

    except Exception as e:

        conn.rollback()

        return f"Delete Error: {e}"

    finally:

        conn.close()


# =========================================================
# STOCK IN
# =========================================================

@app.route(
    "/stock-in",
    methods=["GET", "POST"]
)
def stock_in():

    conn = get_db()
    cursor = conn.cursor()

    if request.method == "POST":

        article_code = request.form.get(
            "article_code",
            ""
        ).strip()

        stock_date = request.form.get(
            "stock_date",
            ""
        ).strip()

        cursor.execute("""
            SELECT article_name
            FROM articles
            WHERE article_code = ?
        """, (article_code,))

        article = cursor.fetchone()

        if not article:

            conn.close()

            return "Article Code not found."

        article_name = article[0]

        stock_sizes = {}

        for size in range(2, 15):

            value = request.form.get(
                f"size_{size}",
                "0"
            ).strip()

            try:
                quantity = int(value)
            except ValueError:
                quantity = 0

            if quantity < 0:
                quantity = 0

            stock_sizes[size] = quantity

        total_qty = sum(
            stock_sizes.values()
        )

        if total_qty <= 0:

            conn.close()

            return "Stock quantity enter karo."

        # -------------------------------------------------
        # UPDATE CURRENT STOCK
        # -------------------------------------------------

        for size in range(2, 15):

            quantity = stock_sizes[size]

            if quantity > 0:

                cursor.execute(
                    f"""
                    UPDATE articles
                    SET size_{size} =
                        size_{size} + ?
                    WHERE article_code = ?
                    """,
                    (
                        quantity,
                        article_code
                    )
                )

        # -------------------------------------------------
        # SAVE STOCK IN HISTORY
        # -------------------------------------------------

        cursor.execute("""
            INSERT INTO stock_entries (
                stock_date,
                article_code,
                article_name,
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14,
                total_qty
            )
            VALUES (
                ?, ?, ?,
                ?, ?, ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?
            )
        """, (
            stock_date,
            article_code,
            article_name,
            stock_sizes[2],
            stock_sizes[3],
            stock_sizes[4],
            stock_sizes[5],
            stock_sizes[6],
            stock_sizes[7],
            stock_sizes[8],
            stock_sizes[9],
            stock_sizes[10],
            stock_sizes[11],
            stock_sizes[12],
            stock_sizes[13],
            stock_sizes[14],
            total_qty
        ))

        conn.commit()
        conn.close()

        return redirect("/")

    cursor.execute("""
        SELECT
            article_code,
            article_name
        FROM articles
        ORDER BY article_code
    """)

    articles = cursor.fetchall()

    conn.close()

    return render_template(
        "stock_in.html",
        articles=articles
    )


# =========================================================
# DISPATCH
# =========================================================

@app.route(
    "/dispatch",
    methods=["GET", "POST"]
)
def dispatch():

    conn = get_db()
    cursor = conn.cursor()

    if request.method == "POST":

        party_name = request.form.get(
            "party_name",
            ""
        ).strip()

        dispatch_date = request.form.get(
            "dispatch_date",
            ""
        ).strip()

        article_code = request.form.get(
            "article_code",
            ""
        ).strip()

        cursor.execute("""
            SELECT
                article_name,
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14
            FROM articles
            WHERE article_code = ?
        """, (article_code,))

        article = cursor.fetchone()

        if not article:

            conn.close()

            return "Article Code not found."

        article_name = article[0]

        current_stock = {}

        for size in range(2, 15):

            current_stock[size] = article[
                size - 1
            ]

        dispatch_sizes = {}

        for size in range(2, 15):

            value = request.form.get(
                f"size_{size}",
                "0"
            ).strip()

            try:
                quantity = int(value)
            except ValueError:
                quantity = 0

            if quantity < 0:
                quantity = 0

            dispatch_sizes[size] = quantity

        # -------------------------------------------------
        # CHECK AVAILABLE STOCK
        # -------------------------------------------------

        for size in range(2, 15):

            if dispatch_sizes[size] > current_stock[size]:

                conn.close()

                return (
                    f"Size {size} me sirf "
                    f"{current_stock[size]} pair available hain."
                )

        total_qty = sum(
            dispatch_sizes.values()
        )

        if total_qty <= 0:

            conn.close()

            return "Dispatch quantity enter karo."

        # -------------------------------------------------
        # REDUCE CURRENT STOCK
        # -------------------------------------------------

        for size in range(2, 15):

            quantity = dispatch_sizes[size]

            if quantity > 0:

                cursor.execute(
                    f"""
                    UPDATE articles
                    SET size_{size} =
                        size_{size} - ?
                    WHERE article_code = ?
                    """,
                    (
                        quantity,
                        article_code
                    )
                )

        # -------------------------------------------------
        # SAVE DISPATCH HISTORY
        # -------------------------------------------------

        cursor.execute("""
            INSERT INTO dispatches (
                dispatch_date,
                party_name,
                article_code,
                article_name,
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14,
                total_qty
            )
            VALUES (
                ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?
            )
        """, (
            dispatch_date,
            party_name,
            article_code,
            article_name,
            dispatch_sizes[2],
            dispatch_sizes[3],
            dispatch_sizes[4],
            dispatch_sizes[5],
            dispatch_sizes[6],
            dispatch_sizes[7],
            dispatch_sizes[8],
            dispatch_sizes[9],
            dispatch_sizes[10],
            dispatch_sizes[11],
            dispatch_sizes[12],
            dispatch_sizes[13],
            dispatch_sizes[14],
            total_qty
        ))

        conn.commit()
        conn.close()

        return redirect("/")

    cursor.execute("""
        SELECT
            article_code,
            article_name
        FROM articles
        ORDER BY article_code
    """)

    articles = cursor.fetchall()

    conn.close()

    return render_template(
        "dispatch.html",
        articles=articles
    )


# =========================================================
# DISPATCH HISTORY
# =========================================================

@app.route("/dispatch-history")
def dispatch_history():

    dispatch_date = request.args.get(
        "dispatch_date",
        ""
    ).strip()

    party_name = request.args.get(
        "party_name",
        ""
    ).strip()

    article_code = request.args.get(
        "article_code",
        ""
    ).strip()

    conn = get_db()
    cursor = conn.cursor()

    query = """
        SELECT
            id,
            dispatch_date,
            party_name,
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14,
            total_qty
        FROM dispatches
        WHERE 1=1
    """

    params = []

    if dispatch_date:

        query += """
            AND dispatch_date = ?
        """

        params.append(dispatch_date)

    if party_name:

        query += """
            AND party_name = ?
        """

        params.append(party_name)

    if article_code:

        query += """
            AND article_code = ?
        """

        params.append(article_code)

    query += """
        ORDER BY id DESC
    """

    cursor.execute(
        query,
        params
    )

    dispatches = cursor.fetchall()

    cursor.execute("""
        SELECT DISTINCT
            party_name
        FROM dispatches
        WHERE party_name IS NOT NULL
          AND party_name != ''
        ORDER BY party_name
    """)

    parties = [
        row[0]
        for row in cursor.fetchall()
    ]

    cursor.execute("""
        SELECT DISTINCT
            article_code,
            article_name
        FROM dispatches
        ORDER BY article_code
    """)

    articles = cursor.fetchall()

    conn.close()

    return render_template(
        "dispatch_history.html",
        dispatches=dispatches,
        parties=parties,
        articles=articles,
        selected_date=dispatch_date,
        selected_party=party_name,
        selected_article=article_code
    )


# =========================================================
# EDIT DISPATCH
# =========================================================

@app.route(
    "/edit-dispatch/<int:dispatch_id>",
    methods=["GET", "POST"]
)
def edit_dispatch(dispatch_id):

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            dispatch_date,
            party_name,
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14,
            total_qty
        FROM dispatches
        WHERE id = ?
    """, (dispatch_id,))

    dispatch_record = cursor.fetchone()

    if not dispatch_record:

        conn.close()

        return "Dispatch record not found."

    if request.method == "POST":

        new_date = request.form.get(
            "dispatch_date",
            ""
        ).strip()

        new_party = request.form.get(
            "party_name",
            ""
        ).strip()

        new_article_code = request.form.get(
            "article_code",
            ""
        ).strip()

        old_article_code = dispatch_record[3]

        old_sizes = {}

        for size in range(2, 15):

            old_sizes[size] = dispatch_record[
                size + 3
            ]

        new_sizes = {}

        for size in range(2, 15):

            value = request.form.get(
                f"size_{size}",
                "0"
            ).strip()

            try:
                quantity = int(value)
            except ValueError:
                quantity = 0

            if quantity < 0:
                quantity = 0

            new_sizes[size] = quantity

        new_total = sum(
            new_sizes.values()
        )

        if new_total <= 0:

            conn.close()

            return "Dispatch quantity enter karo."

        cursor.execute("""
            SELECT
                article_name,
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14
            FROM articles
            WHERE article_code = ?
        """, (new_article_code,))

        new_article = cursor.fetchone()

        if not new_article:

            conn.close()

            return "Article Code not found."

        new_article_name = new_article[0]

        # Restore old dispatch stock

        for size in range(2, 15):

            quantity = old_sizes[size]

            if quantity > 0:

                cursor.execute(
                    f"""
                    UPDATE articles
                    SET size_{size} =
                        size_{size} + ?
                    WHERE article_code = ?
                    """,
                    (
                        quantity,
                        old_article_code
                    )
                )

        # Get available stock

        cursor.execute("""
            SELECT
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14
            FROM articles
            WHERE article_code = ?
        """, (new_article_code,))

        available_article = cursor.fetchone()

        if not available_article:

            conn.rollback()
            conn.close()

            return "Article Code not found."

        # Check stock

        for size in range(2, 15):

            available = available_article[
                size - 2
            ]

            if new_sizes[size] > available:

                conn.rollback()
                conn.close()

                return (
                    f"Size {size} me sirf "
                    f"{available} pair available hain."
                )

        # Reduce new stock

        for size in range(2, 15):

            quantity = new_sizes[size]

            if quantity > 0:

                cursor.execute(
                    f"""
                    UPDATE articles
                    SET size_{size} =
                        size_{size} - ?
                    WHERE article_code = ?
                    """,
                    (
                        quantity,
                        new_article_code
                    )
                )

        # Update history

        cursor.execute("""
            UPDATE dispatches

            SET
                dispatch_date = ?,
                party_name = ?,
                article_code = ?,
                article_name = ?,
                size_2 = ?,
                size_3 = ?,
                size_4 = ?,
                size_5 = ?,
                size_6 = ?,
                size_7 = ?,
                size_8 = ?,
                size_9 = ?,
                size_10 = ?,
                size_11 = ?,
                size_12 = ?,
                size_13 = ?,
                size_14 = ?,
                total_qty = ?

            WHERE id = ?
        """, (
            new_date,
            new_party,
            new_article_code,
            new_article_name,
            new_sizes[2],
            new_sizes[3],
            new_sizes[4],
            new_sizes[5],
            new_sizes[6],
            new_sizes[7],
            new_sizes[8],
            new_sizes[9],
            new_sizes[10],
            new_sizes[11],
            new_sizes[12],
            new_sizes[13],
            new_sizes[14],
            new_total,
            dispatch_id
        ))

        conn.commit()
        conn.close()

        return redirect(
            "/dispatch-history"
        )

    cursor.execute("""
        SELECT
            article_code,
            article_name
        FROM articles
        ORDER BY article_code
    """)

    articles = cursor.fetchall()

    conn.close()

    return render_template(
        "edit_dispatch.html",
        dispatch=dispatch_record,
        articles=articles
    )


# =========================================================
# DELETE DISPATCH
# =========================================================

@app.route(
    "/delete-dispatch/<int:dispatch_id>",
    methods=["POST"]
)
def delete_dispatch(dispatch_id):

    conn = get_db()
    cursor = conn.cursor()

    try:

        cursor.execute("""
            SELECT
                article_code,
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14
            FROM dispatches
            WHERE id = ?
        """, (dispatch_id,))

        record = cursor.fetchone()

        if not record:

            return "Dispatch record not found."

        article_code = record[0]

        for size in range(2, 15):

            quantity = record[
                size - 1
            ]

            if quantity > 0:

                cursor.execute(
                    f"""
                    UPDATE articles
                    SET size_{size} =
                        size_{size} + ?
                    WHERE article_code = ?
                    """,
                    (
                        quantity,
                        article_code
                    )
                )

        cursor.execute("""
            DELETE FROM dispatches
            WHERE id = ?
        """, (dispatch_id,))

        conn.commit()

        return redirect(
            "/dispatch-history"
        )

    except Exception as e:

        conn.rollback()

        return f"Delete Error: {e}"

    finally:

        conn.close()


# =========================================================
# PARTY REPORT
# =========================================================

@app.route("/party-report")
def party_report():

    dispatch_date = request.args.get(
        "dispatch_date",
        ""
    ).strip()

    party_name = request.args.get(
        "party_name",
        ""
    ).strip()

    article_code = request.args.get(
        "article_code",
        ""
    ).strip()

    conn = get_db()
    cursor = conn.cursor()

    query = """
        SELECT
            party_name,
            COUNT(*) AS dispatch_count,
            COALESCE(SUM(size_2), 0),
            COALESCE(SUM(size_3), 0),
            COALESCE(SUM(size_4), 0),
            COALESCE(SUM(size_5), 0),
            COALESCE(SUM(size_6), 0),
            COALESCE(SUM(size_7), 0),
            COALESCE(SUM(size_8), 0),
            COALESCE(SUM(size_9), 0),
            COALESCE(SUM(size_10), 0),
            COALESCE(SUM(size_11), 0),
            COALESCE(SUM(size_12), 0),
            COALESCE(SUM(size_13), 0),
            COALESCE(SUM(size_14), 0),
            COALESCE(SUM(total_qty), 0)
        FROM dispatches
        WHERE 1=1
    """

    params = []

    if dispatch_date:

        query += """
            AND dispatch_date = ?
        """

        params.append(dispatch_date)

    if party_name:

        query += """
            AND party_name = ?
        """

        params.append(party_name)

    if article_code:

        query += """
            AND article_code = ?
        """

        params.append(article_code)

    query += """
        GROUP BY party_name
        ORDER BY party_name
    """

    cursor.execute(
        query,
        params
    )

    party_reports = cursor.fetchall()

    cursor.execute("""
        SELECT DISTINCT
            party_name
        FROM dispatches
        WHERE party_name IS NOT NULL
          AND party_name != ''
        ORDER BY party_name
    """)

    parties = [
        row[0]
        for row in cursor.fetchall()
    ]

    cursor.execute("""
        SELECT DISTINCT
            article_code,
            article_name
        FROM dispatches
        ORDER BY article_code
    """)

    articles = cursor.fetchall()

    conn.close()

    return render_template(
        "party_report.html",
        party_reports=party_reports,
        parties=parties,
        articles=articles,
        selected_date=dispatch_date,
        selected_party=party_name,
        selected_article=article_code
    )


# =========================================================
# STOCK REPORT
# =========================================================

@app.route("/stock-report")
def stock_report():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14
        FROM articles
        ORDER BY article_code
    """)

    stock_reports = cursor.fetchall()

    conn.close()

    return render_template(
        "stock_report.html",
        stock_reports=stock_reports
    )


# =========================================================
# STOCK IN HISTORY
# =========================================================

@app.route("/stock-in-history")
def stock_in_history():

    stock_date = request.args.get(
        "stock_date",
        ""
    ).strip()

    article_code = request.args.get(
        "article_code",
        ""
    ).strip()

    conn = get_db()
    cursor = conn.cursor()

    query = """
        SELECT
            id,
            stock_date,
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14,
            total_qty
        FROM stock_entries
        WHERE 1=1
    """

    params = []

    if stock_date:

        query += """
            AND stock_date = ?
        """

        params.append(stock_date)

    if article_code:

        query += """
            AND article_code = ?
        """

        params.append(article_code)

    query += """
        ORDER BY id DESC
    """

    cursor.execute(
        query,
        params
    )

    stock_entries = cursor.fetchall()

    cursor.execute("""
        SELECT DISTINCT
            article_code,
            article_name
        FROM stock_entries
        ORDER BY article_code
    """)

    articles = cursor.fetchall()

    conn.close()

    return render_template(
        "stock_in_history.html",
        stock_entries=stock_entries,
        articles=articles,
        selected_date=stock_date,
        selected_article=article_code
    )


# =========================================================
# EDIT STOCK IN
# =========================================================

@app.route(
    "/edit-stock-in/<int:stock_id>",
    methods=["GET", "POST"]
)
def edit_stock_in(stock_id):

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            stock_date,
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14,
            total_qty
        FROM stock_entries
        WHERE id = ?
    """, (stock_id,))

    stock_record = cursor.fetchone()

    if not stock_record:

        conn.close()

        return "Stock In record not found."

    if request.method == "POST":

        new_date = request.form.get(
            "stock_date",
            ""
        ).strip()

        new_article_code = request.form.get(
            "article_code",
            ""
        ).strip()

        old_article_code = stock_record[2]

        old_sizes = {}

        for size in range(2, 15):

            old_sizes[size] = stock_record[
                size + 2
            ]

        new_sizes = {}

        for size in range(2, 15):

            value = request.form.get(
                f"size_{size}",
                "0"
            ).strip()

            try:
                quantity = int(value)
            except ValueError:
                quantity = 0

            if quantity < 0:
                quantity = 0

            new_sizes[size] = quantity

        new_total = sum(
            new_sizes.values()
        )

        if new_total <= 0:

            conn.close()

            return "Stock quantity enter karo."

        cursor.execute("""
            SELECT article_name
            FROM articles
            WHERE article_code = ?
        """, (new_article_code,))

        new_article = cursor.fetchone()

        if not new_article:

            conn.close()

            return "Article Code not found."

        new_article_name = new_article[0]

        # -------------------------------------------------
        # OLD ARTICLE STOCK
        # -------------------------------------------------

        cursor.execute("""
            SELECT
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14
            FROM articles
            WHERE article_code = ?
        """, (old_article_code,))

        old_article_stock = cursor.fetchone()

        if not old_article_stock:

            conn.rollback()
            conn.close()

            return "Old Article Code not found."

        # Safety check

        for size in range(2, 15):

            current_quantity = old_article_stock[
                size - 2
            ]

            old_quantity = old_sizes[size]

            if old_quantity > current_quantity:

                conn.rollback()
                conn.close()

                return (
                    f"Stock In Edit nahi ho sakta. "
                    f"Size {size} me current stock "
                    f"{current_quantity} hai, jabki old Stock In "
                    f"{old_quantity} tha."
                )

        # Remove old stock

        for size in range(2, 15):

            quantity = old_sizes[size]

            if quantity > 0:

                cursor.execute(
                    f"""
                    UPDATE articles
                    SET size_{size} =
                        size_{size} - ?
                    WHERE article_code = ?
                    """,
                    (
                        quantity,
                        old_article_code
                    )
                )

        # Add new stock

        for size in range(2, 15):

            quantity = new_sizes[size]

            if quantity > 0:

                cursor.execute(
                    f"""
                    UPDATE articles
                    SET size_{size} =
                        size_{size} + ?
                    WHERE article_code = ?
                    """,
                    (
                        quantity,
                        new_article_code
                    )
                )

        # Update history

        cursor.execute("""
            UPDATE stock_entries

            SET
                stock_date = ?,
                article_code = ?,
                article_name = ?,
                size_2 = ?,
                size_3 = ?,
                size_4 = ?,
                size_5 = ?,
                size_6 = ?,
                size_7 = ?,
                size_8 = ?,
                size_9 = ?,
                size_10 = ?,
                size_11 = ?,
                size_12 = ?,
                size_13 = ?,
                size_14 = ?,
                total_qty = ?

            WHERE id = ?
        """, (
            new_date,
            new_article_code,
            new_article_name,
            new_sizes[2],
            new_sizes[3],
            new_sizes[4],
            new_sizes[5],
            new_sizes[6],
            new_sizes[7],
            new_sizes[8],
            new_sizes[9],
            new_sizes[10],
            new_sizes[11],
            new_sizes[12],
            new_sizes[13],
            new_sizes[14],
            new_total,
            stock_id
        ))

        conn.commit()
        conn.close()

        return redirect(
            "/stock-in-history"
        )

    cursor.execute("""
        SELECT
            article_code,
            article_name
        FROM articles
        ORDER BY article_code
    """)

    articles = cursor.fetchall()

    conn.close()

    return render_template(
        "edit_stock_in.html",
        stock=stock_record,
        articles=articles
    )


# =========================================================
# DELETE STOCK IN
# =========================================================

@app.route(
    "/delete-stock-in/<int:stock_id>",
    methods=["POST"]
)
def delete_stock_in(stock_id):

    conn = get_db()
    cursor = conn.cursor()

    try:

        cursor.execute("""
            SELECT
                article_code,
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14
            FROM stock_entries
            WHERE id = ?
        """, (stock_id,))

        record = cursor.fetchone()

        if not record:

            return "Stock In record not found."

        article_code = record[0]

        cursor.execute("""
            SELECT
                size_2,
                size_3,
                size_4,
                size_5,
                size_6,
                size_7,
                size_8,
                size_9,
                size_10,
                size_11,
                size_12,
                size_13,
                size_14
            FROM articles
            WHERE article_code = ?
        """, (article_code,))

        current_stock = cursor.fetchone()

        if not current_stock:

            return "Article Code not found."

        # Safety check

        for size in range(2, 15):

            stock_quantity = current_stock[
                size - 2
            ]

            stock_in_quantity = record[
                size - 1
            ]

            if stock_in_quantity > stock_quantity:

                return (
                    f"Stock In delete nahi ho sakta. "
                    f"Size {size} me current stock "
                    f"{stock_quantity} hai, jabki is record me "
                    f"{stock_in_quantity} pair hain."
                )

        # Remove stock

        for size in range(2, 15):

            quantity = record[
                size - 1
            ]

            if quantity > 0:

                cursor.execute(
                    f"""
                    UPDATE articles
                    SET size_{size} =
                        size_{size} - ?
                    WHERE article_code = ?
                    """,
                    (
                        quantity,
                        article_code
                    )
                )

        # Delete history

        cursor.execute("""
            DELETE FROM stock_entries
            WHERE id = ?
        """, (stock_id,))

        conn.commit()

        return redirect(
            "/stock-in-history"
        )

    except Exception as e:

        conn.rollback()

        return f"Delete Error: {e}"

    finally:

        conn.close()


# =========================================================
# EXCEL - STOCK REPORT
# =========================================================

@app.route("/export-stock")
def export_stock():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14
        FROM articles
        ORDER BY article_code
    """)

    rows = cursor.fetchall()

    conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stock Report"

    headers = [
        "Article Code",
        "Article Name",
        "Size 2",
        "Size 3",
        "Size 4",
        "Size 5",
        "Size 6",
        "Size 7",
        "Size 8",
        "Size 9",
        "Size 10",
        "Size 11",
        "Size 12",
        "Size 13",
        "Size 14",
        "Total Stock"
    ]

    sheet.append(headers)

    for row in rows:

        total = sum(row[2:15])

        sheet.append(
            list(row) + [total]
        )

    style_excel(sheet)

    file_path = os.path.join(
        os.path.dirname(
            os.path.abspath(__file__)
        ),
        "Stock_Report.xlsx"
    )

    workbook.save(file_path)

    return send_file(
        file_path,
        as_attachment=True,
        download_name="Stock_Report.xlsx"
    )


# =========================================================
# EXCEL - DISPATCH HISTORY
# =========================================================

@app.route("/export-dispatch")
def export_dispatch():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            dispatch_date,
            party_name,
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14,
            total_qty
        FROM dispatches
        ORDER BY id DESC
    """)

    rows = cursor.fetchall()

    conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dispatch History"

    headers = [
        "Date",
        "Party Name",
        "Article Code",
        "Article Name",
        "Size 2",
        "Size 3",
        "Size 4",
        "Size 5",
        "Size 6",
        "Size 7",
        "Size 8",
        "Size 9",
        "Size 10",
        "Size 11",
        "Size 12",
        "Size 13",
        "Size 14",
        "Total"
    ]

    sheet.append(headers)

    for row in rows:

        sheet.append(
            list(row)
        )

    style_excel(sheet)

    file_path = os.path.join(
        os.path.dirname(
            os.path.abspath(__file__)
        ),
        "Dispatch_History.xlsx"
    )

    workbook.save(file_path)

    return send_file(
        file_path,
        as_attachment=True,
        download_name="Dispatch_History.xlsx"
    )


# =========================================================
# EXCEL - PARTY REPORT
# =========================================================

@app.route("/export-party")
def export_party():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            party_name,
            COUNT(*) AS dispatch_count,
            SUM(size_2),
            SUM(size_3),
            SUM(size_4),
            SUM(size_5),
            SUM(size_6),
            SUM(size_7),
            SUM(size_8),
            SUM(size_9),
            SUM(size_10),
            SUM(size_11),
            SUM(size_12),
            SUM(size_13),
            SUM(size_14),
            SUM(total_qty)
        FROM dispatches
        GROUP BY party_name
        ORDER BY party_name
    """)

    rows = cursor.fetchall()

    conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Party Report"

    headers = [
        "Party Name",
        "Dispatch Count",
        "Size 2",
        "Size 3",
        "Size 4",
        "Size 5",
        "Size 6",
        "Size 7",
        "Size 8",
        "Size 9",
        "Size 10",
        "Size 11",
        "Size 12",
        "Size 13",
        "Size 14",
        "Total Pairs"
    ]

    sheet.append(headers)

    for row in rows:

        sheet.append(
            list(row)
        )

    style_excel(sheet)

    file_path = os.path.join(
        os.path.dirname(
            os.path.abspath(__file__)
        ),
        "Party_Report.xlsx"
    )

    workbook.save(file_path)

    return send_file(
        file_path,
        as_attachment=True,
        download_name="Party_Report.xlsx"
    )


# =========================================================
# EXCEL - STOCK IN HISTORY
# =========================================================

@app.route("/export-stock-in")
def export_stock_in():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            stock_date,
            article_code,
            article_name,
            size_2,
            size_3,
            size_4,
            size_5,
            size_6,
            size_7,
            size_8,
            size_9,
            size_10,
            size_11,
            size_12,
            size_13,
            size_14,
            total_qty
        FROM stock_entries
        ORDER BY id DESC
    """)

    rows = cursor.fetchall()

    conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stock In History"

    headers = [
        "Date",
        "Article Code",
        "Article Name",
        "Size 2",
        "Size 3",
        "Size 4",
        "Size 5",
        "Size 6",
        "Size 7",
        "Size 8",
        "Size 9",
        "Size 10",
        "Size 11",
        "Size 12",
        "Size 13",
        "Size 14",
        "Total Stock In"
    ]

    sheet.append(headers)

    for row in rows:

        sheet.append(
            list(row)
        )

    style_excel(sheet)

    file_path = os.path.join(
        os.path.dirname(
            os.path.abspath(__file__)
        ),
        "Stock_In_History.xlsx"
    )

    workbook.save(file_path)

    return send_file(
        file_path,
        as_attachment=True,
        download_name="Stock_In_History.xlsx"
    )


# =========================================================
# EXCEL COMMON STYLE
# =========================================================

def style_excel(sheet):

    for cell in sheet[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for column in sheet.columns:

        max_length = 0

        column_letter = column[0].column_letter

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            25
        )


# =========================================================
# START APPLICATION
# =========================================================

if __name__ == "__main__":

    init_db()

    app.run(
        debug=False
    )